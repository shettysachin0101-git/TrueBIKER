import os
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

def copy_text_frame(src_tf, dest_tf):
    dest_tf.clear()
    dest_tf.word_wrap = src_tf.word_wrap
    dest_tf.margin_left = src_tf.margin_left
    dest_tf.margin_right = src_tf.margin_right
    dest_tf.margin_top = src_tf.margin_top
    dest_tf.margin_bottom = src_tf.margin_bottom
    
    for i, src_p in enumerate(src_tf.paragraphs):
        if i == 0:
            dest_p = dest_tf.paragraphs[0]
        else:
            dest_p = dest_tf.add_paragraph()
        dest_p.alignment = src_p.alignment
        dest_p.space_after = src_p.space_after
        dest_p.space_before = src_p.space_before
        dest_p.line_spacing = src_p.line_spacing
        
        for src_r in src_p.runs:
            dest_r = dest_p.add_run()
            dest_r.text = src_r.text
            if src_r.font:
                dest_r.font.bold = src_r.font.bold
                dest_r.font.italic = src_r.font.italic
                dest_r.font.underline = src_r.font.underline
                dest_r.font.size = src_r.font.size
                dest_r.font.name = src_r.font.name
                try:
                    if src_r.font.color and src_r.font.color.type is not None:
                        dest_r.font.color.rgb = src_r.font.color.rgb
                except Exception:
                    pass

def copy_shape(src_shape, dest_slide, left=None, top=None, width=None, height=None):
    l = left if left is not None else src_shape.left
    t = top if top is not None else src_shape.top
    w = width if width is not None else src_shape.width
    h = height if height is not None else src_shape.height
    
    dest_shape = dest_slide.shapes.add_textbox(l, t, w, h)
    if src_shape.has_text_frame:
        copy_text_frame(src_shape.text_frame, dest_shape.text_frame)
    return dest_shape

def modify_presentation(filename):
    path = os.path.join(r"c:\Users\admin\OneDrive\Desktop\TrueBiker", filename)
    if not os.path.exists(path):
        print(f"File {filename} not found.")
        return
    
    print(f"\nProcessing {filename}...")
    prs = Presentation(path)
    
    # 1. Get original slide 1 and its shapes
    slide_1 = prs.slides[0]
    shapes = list(slide_1.shapes)
    if len(shapes) < 5:
        print(f"Warning: {filename} Slide 1 has less than 5 shapes. Skipping.")
        return
        
    shape_title = shapes[0]
    shape_subtitle = shapes[1]
    shape_intro = shapes[2]
    shape_obj = shapes[3]
    shape_prob = shapes[4]
    
    # 2. Add new slide 2 (using Blank layout, which is layout index 6)
    blank_layout = prs.slide_layouts[6]
    new_slide_2 = prs.slides.add_slide(blank_layout)
    
    # Move the new slide (currently at the end) to index 1 (second slide)
    slide_id_list = prs.slides._sldIdLst
    slides_ids = list(slide_id_list)
    last_slide_id = slides_ids[-1]
    slide_id_list.remove(last_slide_id)
    slide_id_list.insert(1, last_slide_id)
    
    # 3. Populate new Slide 2
    # Title: "INTRODUCTION & OBJECTIVES" or similar
    title_box = new_slide_2.shapes.add_textbox(457200, 365760, 11277295, 731520)
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "INTRODUCTION, OBJECTIVES & PROBLEM STATEMENT"
    p.alignment = PP_ALIGN.LEFT
    
    # Let's style the new Slide 2 Title font based on original title run
    if shape_title.has_text_frame and shape_title.text_frame.paragraphs[0].runs:
        src_run = shape_title.text_frame.paragraphs[0].runs[0]
        r = p.runs[0]
        r.font.bold = True
        r.font.name = src_run.font.name
        # Keep font size smaller since the text is longer
        r.font.size = Pt(32)
        try:
            if src_run.font.color and src_run.font.color.type is not None:
                r.font.color.rgb = src_run.font.color.rgb
        except Exception:
            pass
            
    # Add the three columns copied from original shapes 2, 3, 4
    # Spacing and coordinates match the original columns
    copy_shape(shape_intro, new_slide_2, left=457200, top=1645920, width=3474720, height=4572000)
    copy_shape(shape_obj, new_slide_2, left=4343400, top=1645920, width=3474720, height=4572000)
    copy_shape(shape_prob, new_slide_2, left=8229600, top=1645920, width=3474720, height=4572000)
    
    # 4. Clean up original Slide 1 to make it a dedicated Title Slide
    # Delete shapes 2, 3, 4 (Introduction, Objectives, Problem Statement)
    for s in [shape_intro, shape_obj, shape_prob]:
        el = s.element
        el.getparent().remove(el)
        
    # Center title and subtitle vertically and horizontally
    # Standard height is 6858000 (7.5 in). Center: title top=2200000, subtitle top=3200000
    shape_title.top = 2200000
    shape_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    shape_subtitle.top = 3200000
    shape_subtitle.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # 5. Append "Thank You" slide at the end
    thank_you_slide = prs.slides.add_slide(blank_layout)
    
    # Add title "THANK YOU!"
    ty_title = thank_you_slide.shapes.add_textbox(457200, 2200000, 11277295, 900000)
    tf_ty = ty_title.text_frame
    p_ty = tf_ty.paragraphs[0]
    p_ty.text = "THANK YOU!"
    p_ty.alignment = PP_ALIGN.CENTER
    
    # Style THANK YOU! based on Slide 1 title styling
    if shape_title.has_text_frame and shape_title.text_frame.paragraphs[0].runs:
        src_run = shape_title.text_frame.paragraphs[0].runs[0]
        r = p_ty.runs[0]
        r.font.bold = True
        r.font.name = src_run.font.name
        r.font.size = Pt(54)
        try:
            if src_run.font.color and src_run.font.color.type is not None:
                r.font.color.rgb = src_run.font.color.rgb
        except Exception:
            pass
            
    # Add subtitle "Questions & Discussion"
    ty_sub = thank_you_slide.shapes.add_textbox(457200, 3400000, 11277295, 500000)
    tf_sub = ty_sub.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Questions & Discussion"
    p_sub.alignment = PP_ALIGN.CENTER
    
    if shape_subtitle.has_text_frame and shape_subtitle.text_frame.paragraphs[0].runs:
        src_run = shape_subtitle.text_frame.paragraphs[0].runs[0]
        r = p_sub.runs[0]
        r.font.italic = True
        r.font.name = src_run.font.name
        r.font.size = Pt(24)
        try:
            if src_run.font.color and src_run.font.color.type is not None:
                r.font.color.rgb = src_run.font.color.rgb
        except Exception:
            pass

    # Save presentation
    prs.save(path)
    print(f"Successfully modified and saved {filename}")

modify_presentation("TrueBiker_Project_Synopsis.pptx")
modify_presentation("TrueBiker_Project_Synopsis1.pptx")
